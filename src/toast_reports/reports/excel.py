"""Excel workbook renderer.

Layout:
  * "Summary" tab  — one row per (week, location) with every metric, plus a
    company total row per week.
  * One tab per location — that location's week-over-week trend.
  * "Revenue Centers" tab — orders per day per revenue center per location, as
    tidy rows (one row = location x day x revenue center) so it pivots cleanly.

Currency/percent number formats are applied so the file is ready to read, not
just a dump of numbers.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..aggregate import CompanyWeek, WeeklyMetrics, group_by_week

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FONT = Font(bold=True)
_MONEY = "$#,##0.00"
_PCT = "0.0%"
_NUM = "#,##0.0"

_COLUMNS = [
    ("Week Of", "week_start", None),
    ("Location", "location_name", None),
    ("Net Sales", "net_sales", _MONEY),
    ("Gross Sales", "gross_sales", _MONEY),
    ("Tax", "tax", _MONEY),
    ("Tips", "tips", _MONEY),
    ("Transactions", "transaction_count", "#,##0"),
    ("Guests", "guest_count", "#,##0"),
    ("Avg Check", "avg_check", _MONEY),
    ("Labor Hours", "labor_hours", _NUM),
    ("Overtime Hrs", "overtime_hours", _NUM),
    ("Labor Cost", "labor_cost", _MONEY),
    ("Labor %", "labor_pct", _PCT),
    ("Sales / Labor Hr", "sales_per_labor_hour", _MONEY),
]


def _value(m: WeeklyMetrics, attr: str):
    if attr == "week_start":
        return m.week_start.isoformat()
    return getattr(m, attr)


def _write_header(ws: Worksheet) -> None:
    for col, (label, _attr, _fmt) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _write_row(ws: Worksheet, row: int, m: WeeklyMetrics, bold: bool = False) -> None:
    for col, (_label, attr, fmt) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=_value(m, attr))
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = _TOTAL_FONT


def _autosize(ws: Worksheet) -> None:
    for col in range(1, len(_COLUMNS) + 1):
        width = max(
            (len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 12), 22)


def _write_grid(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    """Write an arbitrary header + rows block (used by tabs that aren't the
    WeeklyMetrics column layout), then size the columns to fit."""
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A2"
    for col in range(1, len(headers) + 1):
        width = max(
            (len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 12), 28)


def _revenue_center_rows(revenue_centers: dict) -> list[list]:
    """Flatten {location: {days, centers, counts}} into tidy rows. Zero-order
    (location, day, center) combinations are written as 0 rather than skipped,
    so a pivot over the block has no gaps."""
    rows: list[list] = []
    for name in sorted(revenue_centers):
        rc = revenue_centers[name]
        for day in rc["days"]:
            for center in rc["centers"]:
                rows.append([name, day, center, int(rc["counts"].get(f"{day}|{center}", 0))])
    return rows


def _company_total_row(week: CompanyWeek) -> WeeklyMetrics:
    total = WeeklyMetrics(
        location_guid="ALL",
        location_name="All Locations",
        week_start=week.week_start,
    )
    for r in week.rows:
        total.net_sales += r.net_sales
        total.gross_sales += r.gross_sales
        total.tax += r.tax
        total.tips += r.tips
        total.transaction_count += r.transaction_count
        total.check_count += r.check_count
        total.guest_count += r.guest_count
        total.regular_hours += r.regular_hours
        total.overtime_hours += r.overtime_hours
        total.labor_cost += r.labor_cost
    return total


def render_workbook(
    metrics: list[WeeklyMetrics],
    out_path: str | Path,
    title: str,
    revenue_centers: dict | None = None,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_header(summary)

    row = 2
    for week in group_by_week(metrics):
        for m in week.rows:
            _write_row(summary, row, m)
            row += 1
        _write_row(summary, row, _company_total_row(week), bold=True)
        row += 1
    _autosize(summary)

    # Per-location tabs.
    by_loc: dict[str, list[WeeklyMetrics]] = {}
    for m in metrics:
        by_loc.setdefault(m.location_name, []).append(m)

    for name, rows in by_loc.items():
        ws = wb.create_sheet(title=_safe_sheet_name(name))
        _write_header(ws)
        for i, m in enumerate(sorted(rows, key=lambda x: x.week_start), start=2):
            _write_row(ws, i, m)
        _autosize(ws)

    # Orders per day per revenue center — only when some location uses them.
    if revenue_centers:
        ws = wb.create_sheet(title="Revenue Centers")
        _write_grid(
            ws,
            ["Location", "Business Date", "Revenue Center", "Orders"],
            _revenue_center_rows(revenue_centers),
        )

    wb.save(out_path)
    return out_path


def _safe_sheet_name(name: str) -> str:
    # Excel sheet names: max 31 chars, no []:*?/\
    for ch in "[]:*?/\\":
        name = name.replace(ch, "-")
    return name[:31]
